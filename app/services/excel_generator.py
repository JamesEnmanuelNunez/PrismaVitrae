import io
from collections.abc import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(name="Cambria", size=14, bold=True, color="000000")
HEADER_FILL = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
TITLE_FONT = Font(name="Cambria", size=28, bold=False)
SUBTITLE_FONT = Font(name="Cambria", size=16, bold=False)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_title_rows(ws, title: str, subtitle: str, max_col: int) -> int:
    title_row = 1
    subtitle_row = 2
    header_row = 3

    ws.merge_cells(
        start_row=title_row, start_column=1, end_row=title_row, end_column=max_col
    )
    cell = ws.cell(row=title_row, column=1, value=title)
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(
        start_row=subtitle_row, start_column=1, end_row=subtitle_row, end_column=max_col
    )
    cell = ws.cell(row=subtitle_row, column=1, value=subtitle)
    cell.font = SUBTITLE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

    return header_row


def _style_header(ws, headers: list[str], header_row: int = 3) -> None:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN


def _auto_width(ws) -> None:
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)


def _get(d, key, default=""):
    val = d.get(key) if isinstance(d, dict) else getattr(d, key, default)
    return val if val is not None else default


def _date(value):
    return value.strftime("%Y-%m-%d") if hasattr(value, "strftime") else value


def _generate_sheet(
    *,
    sheet_title: str,
    subtitle: str,
    headers: list[str],
    rows: list[dict],
    row_builder: Callable[[dict], list],
) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_row = _add_title_rows(
        ws, "DIRECCIÓN DE RECURSOS HUMANOS", subtitle, len(headers)
    )
    _style_header(ws, headers, header_row)

    for row_idx, item in enumerate(rows, header_row + 1):
        for col_idx, value in enumerate(row_builder(item), 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    _auto_width(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


CANDIDATOS_HEADERS = [
    "No.", "Reg/Dist", "Cédula", "Nombre Completo", "Sexo",
    "Cargo Solicitado", "Escolaridad", "En Sustitución De",
    "Cédula No.", "Fecha Ingreso", "Centro", "Referido Por", "Teléfono",
]

PROPUESTAS_HEADERS = CANDIDATOS_HEADERS

REAJUSTES_HEADERS = [
    "Cédula", "Nombre Completo", "Grupo Ocupacional", "Cargo",
    "Salario Actual", "Salario Solicitado", "Diferencia",
    "% Incremento", "Observación", "Fecha Efectividad",
]

NO_PROCEDEN_HEADERS = [
    "No.", "Reg/Dist", "Cédula", "Nombre Completo", "Sexo",
    "Cargo Solicitado", "Salario Solicitado", "Escolaridad",
    "Observación", "Referido Por", "Teléfono",
]


def _candidato_row(c: dict) -> list:
    dc = c.get("datos_crudos", {}) or {}
    return [
        dc.get("no", ""),
        dc.get("reg_dist", ""),
        dc.get("cedula", ""),
        _get(c, "nombre"),
        dc.get("sexo", ""),
        dc.get("cargo_solicitado", ""),
        dc.get("escolaridad", "") or _get(c, "educacion"),
        dc.get("en_sustitucion_de", ""),
        dc.get("cedula_no", ""),
        dc.get("fecha_ingreso", ""),
        dc.get("centro", ""),
        dc.get("referido_por", ""),
        _get(c, "telefono"),
    ]


def _propuesta_row(p: dict) -> list:
    return [
        _get(p, "no"),
        _get(p, "reg_dist"),
        _get(p, "cedula"),
        _get(p, "nombre_completo"),
        _get(p, "sexo"),
        _get(p, "cargo_solicitado"),
        _get(p, "escolaridad"),
        _get(p, "en_sustitucion_de"),
        _get(p, "cedula_no"),
        _date(_get(p, "fecha_ingreso", "")),
        _get(p, "centro"),
        _get(p, "referido_por"),
        _get(p, "telefono"),
    ]


def _reajuste_row(r: dict) -> list:
    return [
        _get(r, "cedula"),
        _get(r, "nombre_completo"),
        _get(r, "grupo_ocupacional"),
        _get(r, "cargo"),
        _get(r, "salario_actual", 0),
        _get(r, "salario_solicitado", 0),
        _get(r, "diferencia_salarial", 0),
        _get(r, "porcentaje_incremento", 0),
        _get(r, "observacion"),
        _date(_get(r, "fecha_efectividad", "")),
    ]


def _no_procede_row(np: dict) -> list:
    return [
        _get(np, "no"),
        _get(np, "reg_dist"),
        _get(np, "cedula"),
        _get(np, "nombre_completo"),
        _get(np, "sexo"),
        _get(np, "cargo_solicitado"),
        _get(np, "salario_solicitado", 0),
        _get(np, "escolaridad"),
        _get(np, "observacion"),
        _get(np, "referido_por"),
        _get(np, "telefono"),
    ]


def generate_candidatos_excel(candidatos: list[dict]) -> io.BytesIO:
    return _generate_sheet(
        sheet_title="Candidatos",
        subtitle="CANDIDATOS",
        headers=CANDIDATOS_HEADERS,
        rows=candidatos,
        row_builder=_candidato_row,
    )


def generate_propuestas_excel(propuestas: list[dict]) -> io.BytesIO:
    return _generate_sheet(
        sheet_title="Propuestas",
        subtitle="PROPUESTAS REGIONAL 2025",
        headers=PROPUESTAS_HEADERS,
        rows=propuestas,
        row_builder=_propuesta_row,
    )


def generate_reajustes_excel(reajustes: list[dict]) -> io.BytesIO:
    return _generate_sheet(
        sheet_title="Reajustes",
        subtitle="REAJUSTE SALARIAL",
        headers=REAJUSTES_HEADERS,
        rows=reajustes,
        row_builder=_reajuste_row,
    )


def generate_no_proceden_excel(no_proceden: list[dict]) -> io.BytesIO:
    return _generate_sheet(
        sheet_title="No Proceden",
        subtitle="NO PROCEDEN 2025",
        headers=NO_PROCEDEN_HEADERS,
        rows=no_proceden,
        row_builder=_no_procede_row,
    )
