import uuid
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

CANDIDATO_DATA = {
    "nombre": "Export Test",
    "telefono": "809-555-0000",
    "email": "export@test.com",
    "habilidades": ["Excel", "Data"],
    "experiencia_anios": 2.0,
    "educacion": "Técnico",
}

PROPUESTA_DATA = {
    "no": 1,
    "cedula": "001-1234567-8",
    "nombre_completo": "Propuesta Export",
    "sexo": "M",
    "cargo_solicitado": "Analista",
}

REAJUSTE_DATA = {
    "cedula": "001-9876543-2",
    "nombre_completo": "Reajuste Export",
    "grupo_ocupacional": "Profesional",
    "cargo": "Ingeniero",
    "salario_actual": 35000.0,
    "salario_solicitado": 42000.0,
}

NO_PROCEDER_DATA = {
    "no": 1,
    "cedula": "001-5555555-5",
    "nombre_completo": "No Proceder Export",
    "sexo": "F",
    "cargo_solicitado": "Gerente",
    "salario_solicitado": 70000.0,
    "escolaridad": "Maestría",
}


def _unique(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}"


def _export_values(response) -> tuple[str, list]:
    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    return ws.title, values


def _track(client, cleanup, table, payload):
    response = client.post(f"/api/{table}/", json=payload)
    assert response.status_code == 201
    row = response.json()
    cleanup.table(table, row["id"])
    return row


def test_export_candidatos(client: TestClient, cleanup):
    name = _unique("Export Test")
    _track(client, cleanup, "candidatos", {**CANDIDATO_DATA, "nombre": name})

    response = client.get("/api/exportar-excel/candidatos")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    title, values = _export_values(response)
    assert title == "Candidatos"
    assert name in values


def test_export_propuestas(client: TestClient, cleanup):
    name = _unique("Propuesta Export")
    _track(client, cleanup, "propuestas", {**PROPUESTA_DATA, "nombre_completo": name})

    response = client.get("/api/exportar-excel/propuestas")
    assert response.status_code == 200

    title, values = _export_values(response)
    assert title == "Propuestas"
    assert name in values


def test_export_reajustes(client: TestClient, cleanup):
    name = _unique("Reajuste Export")
    _track(client, cleanup, "reajustes", {**REAJUSTE_DATA, "nombre_completo": name})

    response = client.get("/api/exportar-excel/reajustes")
    assert response.status_code == 200

    title, values = _export_values(response)
    assert title == "Reajustes"
    assert name in values


def test_export_no_proceden(client: TestClient, cleanup):
    name = _unique("No Proceder Export")
    _track(
        client,
        cleanup,
        "no-proceden",
        {**NO_PROCEDER_DATA, "nombre_completo": name},
    )

    response = client.get("/api/exportar-excel/no_proceden")
    assert response.status_code == 200

    title, values = _export_values(response)
    assert title == "No Proceden"
    assert name in values


def test_export_invalid_table(client: TestClient):
    response = client.get("/api/exportar-excel/tabla_inexistente")
    assert response.status_code == 404
    assert "no encontrada" in response.json()["detail"]